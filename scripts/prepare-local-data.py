#!/usr/bin/env python3
"""Read the four supplied XLSX inputs and prepare a private local preview artifact.

Workbook values travel to Node through stdin, never through an intermediate file.
Only aggregate counts are printed. Originals are opened read-only and never saved.
"""
import argparse
import hashlib
import io
import json
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SOURCES = [
    ('Senior', 12, 'student_responses', 'Copy of Project_ Pomfret Ideation for Seniors - for AI sorting.xlsx'),
    ('Junior', 11, 'student_responses', 'Copy of Project_ Pomfret Ideation for Juniors for AI sorting.xlsx'),
    ('Sophomore', 10, 'student_responses', 'Copy of Project_ Pomfret Ideation for Sophomores for AI sorting.xlsx'),
    ('Projects', None, 'project_catalog', 'Copy of Project_ Pomfret 2026-27 Project List for Sorting.xlsx'),
]


def encode(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def digest(path):
    checksum = hashlib.sha256()
    with path.open('rb') as source:
        for block in iter(lambda: source.read(1024 * 1024), b''):
            checksum.update(block)
    return checksum.hexdigest()


def background_of(cell):
    # openpyxl read-only worksheets represent absent cells with fill=None.
    fill = getattr(cell, 'fill', None)
    color = getattr(fill, 'fgColor', None)
    rgb = getattr(color, 'rgb', None)
    if getattr(color, 'type', None) == 'rgb' and isinstance(rgb, str) and len(rgb) >= 6:
        return '#' + rgb[-6:]
    return '#ffffff'


def read_source(source_dir, name, grade, role, filename):
    from openpyxl import load_workbook
    path = source_dir / filename
    # Reading the bytes once makes the parsed data correspond to this precise hash.
    contents = path.read_bytes()
    sha256 = hashlib.sha256(contents).hexdigest()
    book = load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
    try:
        if len(book.worksheets) != 1:
            raise ValueError('Unexpected source sheet count')
        sheet = book.worksheets[0]
        values, backgrounds = [], []
        for cells in sheet.iter_rows():
            values.append([encode(cell.value) for cell in cells])
            backgrounds.append([background_of(cell) for cell in cells])
        result = {
            'name': name, 'grade': grade, 'role': role,
            'values': values, 'backgrounds': backgrounds,
            'sourceFile': filename, 'originalSheet': sheet.title, 'sha256': sha256,
        }
    finally:
        book.close()
    if digest(path) != sha256:
        raise ValueError('A source changed while it was being read')
    return result


def main():
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument('--source-dir', type=Path, required=True)
    cli.add_argument('--output-dir', type=Path, default=ROOT.parent / 'case-f-private')
    cli.add_argument('--node', default='node')
    cli.add_argument('--handoff', type=Path,
                     help='Optional original handoff; only filename and checksum are retained.')
    args = cli.parse_args()
    stage = 'read_sources'
    child_diagnostic = None
    try:
        loaded = [read_source(args.source_dir, *source) for source in SOURCES]
        stage = 'read_references'
        reference_paths = [
            (args.source_dir / 'Copy of 2026 Project_ Pomfret - STUDENT ROSTERS.xlsx', 'historical_output_layout'),
            (args.source_dir / 'Case F discovery meeting transcript (Mr. Litowitz, Oct 20).docx', 'requirements'),
        ]
        if args.handoff:
            reference_paths.append((args.handoff, 'later_defaults'))
        references = [(path, {'name': path.name, 'role': role, 'sha256': digest(path)})
                      for path, role in reference_paths if path.is_file()]
        # Verify every original after all four reads, before creating any output.
        stage = 'verify_sources'
        unchanged = all(digest(args.source_dir / source['sourceFile']) == source['sha256'] for source in loaded)
        unchanged = unchanged and all(digest(path) == metadata['sha256'] for path, metadata in references)
        if not unchanged:
            raise ValueError('Source verification failed')
        payload = {'schemaVersion': 1, 'tables': loaded[:3], 'projectTable': loaded[3],
                   'referenceFiles': [metadata for _, metadata in references], 'sourceIntegrity': {'unchanged': True}}
        stage = 'node_prepare'
        result = subprocess.run(
            [str(args.node), str(ROOT / 'scripts/prepare-local-data.cjs'), '--output-dir', str(args.output_dir)],
            input=json.dumps(payload, ensure_ascii=False, allow_nan=False), text=True, capture_output=True, check=False,
        )
        if result.returncode:
            try:
                candidate = json.loads(result.stderr)
                allowed_stages = {'arguments', 'read_input', 'parse_input', 'prepare_data', 'write_output', 'receipt'}
                allowed_types = {'Error', 'TypeError', 'RangeError', 'SyntaxError'}
                allowed_codes = {'INVALID_SOURCE', 'INVALID_SOURCE_FILE', 'INVALID_SOURCE_SHEET', 'INVALID_SOURCE_METADATA',
                                 'INVALID_PAYLOAD', 'UNVERIFIED_SOURCES', 'MISSING_GRADE', 'DUPLICATE_GRADE_TAB',
                                 'INVALID_PROJECTS', 'INVALID_STUDENTS', 'UNRESOLVED_SOURCE', 'INVALID_REFERENCES',
                                 'INVALID_REFERENCE_FILE', 'INVALID_REFERENCE_METADATA', 'OUTPUT_INSIDE_CODE',
                                 'INVALID_OUTPUT_DIRECTORY', 'INVALID_OUTPUT_FILE', 'REFUSE_UNRELATED_FILE',
                                 'INVALID_ARGUMENTS', 'INPUT_TOO_LARGE', 'PARSER_ERROR', 'UNEXPECTED_ERROR',
                                 'EACCES', 'EPERM', 'ENOENT', 'EEXIST', 'ENOTDIR', 'EISDIR'}
                if (set(candidate) == {'status', 'stage', 'code', 'type'} and candidate['status'] == 'failed'
                        and candidate['stage'] in allowed_stages and candidate['type'] in allowed_types
                        and candidate['code'] in allowed_codes):
                    child_diagnostic = candidate
            except Exception:
                pass
            raise ValueError('Local preparation failed')
        # The child emits a deliberately small receipt; never relay unknown output.
        stage = 'verify_receipt'
        receipt = json.loads(result.stdout)
        safe_keys = {'status', 'schemaVersion', 'sourceFiles', 'students', 'eligible', 'excluded', 'review', 'projects', 'originalsUnchanged', 'artifact'}
        if set(receipt) != safe_keys or receipt.get('status') != 'written':
            raise ValueError('Unexpected preparation receipt')
        print(json.dumps(receipt))
        return 0
    except Exception as error:
        # Exceptions from XLSX/parser dependencies can contain source values or names.
        # Keep the command output private even when the input is malformed.
        diagnostic = {'status': 'failed', 'stage': stage, 'type': type(error).__name__, 'code': 'LOCAL_IMPORT_FAILED'}
        if child_diagnostic:
            diagnostic['child'] = child_diagnostic
        print(json.dumps(diagnostic), file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
